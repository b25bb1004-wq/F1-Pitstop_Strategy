"""Builds presentation/Pitwall.pptx: the Pitwall story (how it works, how it failed, how it evolved)
in the app's F1 broadcast theme. Every number comes from reports/metrics.json and the project notes.

    pip install python-pptx pillow
    python presentation/build_deck.py

Fonts (OFL, in presentation/fonts): Barlow Condensed Bold / Bold Italic, Barlow, JetBrains Mono.
Install them before opening the deck, or PowerPoint/Keynote will substitute.
"""
import os

from lxml import etree
from PIL import Image
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION, XL_LEGEND_POSITION, XL_MARKER_STYLE, XL_TICK_MARK
from pptx.enum.dml import MSO_LINE
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

HERE = os.path.dirname(os.path.abspath(__file__))
OUT = os.path.join(HERE, "Pitwall.pptx")
A = lambda f: os.path.join(HERE, "assets", f)

# theme (same tokens as web/src/styles.css)
BG, PANEL, PANEL2, LINE, HAIR = "0B0B10", "15151E", "1C1C27", "2A2A36", "1E1E28"
RED, TEXT, TEXT2, TEXT3 = "E10600", "F5F5F7", "A9A9B8", "6E6E7E"
BLUE, GREEN, PURPLE, YELLOW = "2A8BDB", "2BA84A", "A855F7", "FFD12E"
SOFT, MED, HARD = "FF3B3B", "FFD12E", "F2F2F2"
COND, BODY, MONO = "Barlow Condensed", "Barlow", "JetBrains Mono"
TOTAL = 20

prs = Presentation()
prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
BLANK = prs.slide_layouts[6]
rgb = lambda h: RGBColor.from_string(h)


# ---------- primitives ----------
def style(run, size=14, font=BODY, color=TEXT, bold=False, italic=False, spacing=None):
    f = run.font
    f.size, f.name, f.bold, f.italic = Pt(size), font, bold, italic
    f.color.rgb = rgb(color)
    if spacing:
        run._r.get_or_add_rPr().set("spc", str(int(spacing * 100)))


def tb(s, x, y, w, h, content, align="l", anchor="t", line=None, after=0, **st):
    """content: str (\\n splits paragraphs), [(text, style)] for one paragraph, or [[(text, style)], ...]."""
    box = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = {"t": MSO_ANCHOR.TOP, "m": MSO_ANCHOR.MIDDLE, "b": MSO_ANCHOR.BOTTOM}[anchor]
    if isinstance(content, str):
        paras = [[(t, {})] for t in content.split("\n")]
    elif content and isinstance(content[0], tuple):
        paras = [content]
    else:
        paras = content
    for i, runs in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = {"l": PP_ALIGN.LEFT, "c": PP_ALIGN.CENTER, "r": PP_ALIGN.RIGHT}[align]
        if line:
            p.line_spacing = line
        if after:
            p.space_after = Pt(after)
        for text, over in runs:
            r = p.add_run()
            r.text = text
            style(r, **{**st, **over})
    return box


def rect(s, x, y, w, h, fill=None, line=None, lw=0.75, shape=MSO_SHAPE.RECTANGLE, dash=False):
    sh = s.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    if fill:
        sh.fill.solid()
        sh.fill.fore_color.rgb = rgb(fill)
    else:
        sh.fill.background()
    if line:
        sh.line.color.rgb = rgb(line)
        sh.line.width = Pt(lw)
        if dash:
            sh.line.dash_style = MSO_LINE.DASH
    else:
        sh.line.fill.background()
    sh.shadow.inherit = False
    return sh


def tag(s, x, y, w, h, text, fill=RED, color=TEXT, size=12, skew=0.25, line=None):
    sh = rect(s, x, y, w, h, fill, line=line, shape=MSO_SHAPE.PARALLELOGRAM)
    sh.adjustments[0] = skew
    tf = sh.text_frame
    tf.margin_left = tf.margin_right = Inches(0.04)
    tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    tf.word_wrap = False
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    style(p.add_run(), size=size, font=COND, bold=True, italic=True, color=color, spacing=1)
    p.runs[0].text = text.upper()
    return sh


def line(s, x1, y1, x2, y2, color=TEXT3, w=1.0, dash=False, head=False):
    c = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    c.line.color.rgb = rgb(color)
    c.line.width = Pt(w)
    if dash:
        c.line.dash_style = MSO_LINE.DASH
    if head:
        te = etree.SubElement(c.line._get_or_add_ln(), qn("a:tailEnd"))
        te.set("type", "triangle")
        te.set("w", "med")
        te.set("len", "med")
    return c


def poly(s, pts, color, w=2.0):
    fb = s.shapes.build_freeform(Inches(pts[0][0]), Inches(pts[0][1]), scale=1.0)
    fb.add_line_segments([(Inches(x), Inches(y)) for x, y in pts[1:]], close=False)
    sh = fb.convert_to_shape()
    sh.fill.background()
    sh.line.color.rgb = rgb(color)
    sh.line.width = Pt(w)
    sh.shadow.inherit = False
    return sh


def pic(s, f, x, y, w, h=None, border=LINE):
    """Place an image; with h given, crop to fill the box (no stretching)."""
    if h is None:
        p = s.shapes.add_picture(A(f), Inches(x), Inches(y), Inches(w))
    else:
        iw, ih = Image.open(A(f)).size
        p = s.shapes.add_picture(A(f), Inches(x), Inches(y), Inches(w), Inches(h))
        rb, ri = w / h, iw / ih
        if ri > rb:
            c = (1 - rb / ri) / 2
            p.crop_left = p.crop_right = c
        else:
            c = (1 - ri / rb) / 2
            p.crop_top = p.crop_bottom = c
    if border:
        p.line.color.rgb = rgb(border)
        p.line.width = Pt(0.75)
    return p


def dot(s, cx, cy, d, fill, line_c=None):
    return rect(s, cx - d / 2, cy - d / 2, d, d, fill, line=line_c, lw=1.5, shape=MSO_SHAPE.OVAL)


# ---------- slide furniture ----------
def new_slide(n, section, bg="bg_grid.jpg", note=""):
    s = prs.slides.add_slide(BLANK)
    s.background.fill.solid()
    s.background.fill.fore_color.rgb = rgb(BG)
    s.shapes.add_picture(A(bg), 0, 0, prs.slide_width, prs.slide_height)
    # footer: brand, section, lap counter and a race-progress line
    tb(s, 0.6, 6.98, 6, 0.3, [("//", {"color": RED}), (" PITWALL", {}), (f"   ·   {section.upper()}", {"color": TEXT3, "italic": False, "size": 11})],
       size=13, font=COND, bold=True, italic=True, spacing=1.5)
    tb(s, 9.73, 6.99, 3.0, 0.3, [("LAP ", {"color": TEXT3}), (f"{n:02d}", {"color": TEXT, "bold": True}), (f" / {TOTAL}", {"color": TEXT3})],
       align="r", size=11, font=MONO)
    rect(s, 0, 7.42, 13.333, 0.04, HAIR)
    rect(s, 0, 7.42, 13.333 * n / TOTAL, 0.04, RED)
    if note:
        s.notes_slide.notes_text_frame.text = note
    return s


def kicker(s, text, x=0.6, y=0.55):
    rect(s, x, y + 0.09, 0.34, 0.04, RED)
    tb(s, x + 0.48, y, 10, 0.3, text.upper(), size=12, font=COND, bold=True, color=RED, spacing=3)


def title(s, parts, x=0.6, y=0.82, w=12.2, size=46, h=1.0):
    tb(s, x, y, w, h, [(t.upper(), {"color": c}) for t, c in parts], size=size, font=COND, bold=True, italic=True, line=0.9)


def label(s, x, y, text, color=TEXT3, w=4, size=10.5, align="l"):
    tb(s, x, y, w, 0.25, text.upper(), size=size, font=COND, bold=True, color=color, spacing=2, align=align)


def chart_clean(chart):
    cs = chart._chartSpace
    for parent, anchor in ((cs, cs.find(qn("c:chart"))), (cs.find(".//" + qn("c:plotArea")), None)):
        sp = parent.find(qn("c:spPr"))
        if sp is None:
            sp = etree.SubElement(parent, qn("c:spPr"))
            if anchor is not None:
                anchor.addnext(sp)
        for ch in list(sp):
            sp.remove(ch)
        etree.SubElement(sp, qn("a:noFill"))
        ln = etree.SubElement(sp, qn("a:ln"))
        etree.SubElement(ln, qn("a:noFill"))
    chart.font.name, chart.font.size = MONO, Pt(11)
    chart.font.color.rgb = rgb(TEXT2)


def hbar_chart(s, x, y, w, h, cats, vals, colors, fmt, vmax):
    cd = CategoryChartData()
    cd.categories = cats
    cd.add_series("v", vals)
    ch = s.shapes.add_chart(XL_CHART_TYPE.BAR_CLUSTERED, Inches(x), Inches(y), Inches(w), Inches(h), cd).chart
    chart_clean(ch)
    ch.has_title, ch.has_legend = False, False
    pl = ch.plots[0]
    pl.gap_width, pl.vary_by_categories = 55, False
    ser = pl.series[0]
    for i, c in enumerate(colors):
        pt = ser.points[i]
        pt.format.fill.solid()
        pt.format.fill.fore_color.rgb = rgb(c)
        pt.format.line.fill.background()
    pl.has_data_labels = True
    dl = pl.data_labels
    dl.number_format, dl.number_format_is_linked = fmt, False
    dl.position = XL_LABEL_POSITION.OUTSIDE_END
    dl.font.size, dl.font.name, dl.font.bold = Pt(13), MONO, True
    dl.font.color.rgb = rgb(TEXT)
    va, ca = ch.value_axis, ch.category_axis
    va.visible, va.has_major_gridlines = False, False
    va.minimum_scale, va.maximum_scale = 0, vmax
    ca.reverse_order = True
    ca.major_tick_mark = XL_TICK_MARK.NONE
    ca.format.line.color.rgb = rgb(LINE)
    ca.tick_labels.font.size, ca.tick_labels.font.name = Pt(13), COND
    ca.tick_labels.font.bold = True
    ca.tick_labels.font.color.rgb = rgb(TEXT)
    return ch


# =====================================================================================
# 01 · LIGHTS OUT
s = new_slide(1, "Lights out", bg="bg_title_car.jpg", note=(
    "Pitwall answers the question a Formula 1 pit wall asks every lap: box now, or stay out? "
    "It predicts lap times, models how tyres wear, measures what a stop costs and makes the call, "
    "trained on every race from 2021 to 2025. This is the story of how it works, how it failed, and how it evolved."))
for i in range(5):  # five start lights, all lit
    x = 8.55 + i * 0.86
    rect(s, x, 0.55, 0.62, 1.32, "121218", line=LINE, shape=MSO_SHAPE.ROUNDED_RECTANGLE).adjustments[0] = 0.18
    for j in range(2):
        dot(s, x + 0.31, 0.88 + j * 0.66, 0.42, RED if j else "3A0A08")
kicker(s, "F1 strategy engine  ·  FastF1 data 2021-2025", y=0.62)
tb(s, 0.55, 1.0, 7.6, 1.8, [("PIT", {}), ("WALL", {"color": RED})], size=118, font=COND, bold=True, italic=True, line=0.85)
tb(s, 0.6, 2.85, 4.4, 1.3, "Predicting lap times, tyre wear and the pit call from 125,046 laps of Formula 1.", size=20, color=TEXT2, line=1.1)
x = 0.6
for t, w in (("Lap time", 1.15), ("Tyre wear", 1.25), ("Box or stay out", 1.75), ("What if", 1.05)):
    tag(s, x, 4.3, w, 0.38, t, fill=PANEL2, color=TEXT, size=13, line=LINE)
    x += w + 0.08
tb(s, 0.6, 6.25, 6, 0.3, [("ARNAV YADAV", {"color": TEXT, "bold": True}), ("   ·   2026-10-06", {})], size=12, font=MONO, color=TEXT2)

# 02 · THE QUESTION
s = new_slide(2, "The question", note=(
    "Every lap the tyres get a little slower, and a stop costs about 22.6 seconds. The pit wall's job is finding the lap "
    "where fresh tyres win back more than the stop costs, while a Safety Car can make a stop nearly free. "
    "So the project has three jobs: predict the lap time, model the tyre, decide the call."))
kicker(s, "The question")
title(s, [("Every lap, ", TEXT), ("one question.", RED)])
tb(s, 0.6, 2.05, 5.6, 0.8, [("“Box, or stay out?”", {})], size=40, font=COND, bold=True, italic=True)
tb(s, 0.6, 2.95, 5.4, 1.7, "A pit stop costs about 22.6 seconds. Old tyres cost a little more every lap. The pit wall has to find the lap where fresh "
   "tyres win back more than the stop costs, knowing a Safety Car could make the stop almost free.", size=15, color=TEXT2, line=1.2)
for i, (n, a, b) in enumerate((("01", "Predict", "the lap time"), ("02", "Model", "the tyre"), ("03", "Decide", "the call"))):
    x = 0.6 + i * 1.86
    rect(s, x, 4.85, 1.74, 1.45, PANEL, line=LINE)
    rect(s, x, 4.85, 1.74, 0.05, RED if i == 2 else LINE)
    tb(s, x + 0.18, 5.03, 1.4, 0.3, n, size=11, font=MONO, color=RED)
    tb(s, x + 0.18, 5.3, 1.5, 0.45, a.upper(), size=24, font=COND, bold=True, italic=True)
    tb(s, x + 0.18, 5.78, 1.5, 0.3, b, size=13, color=TEXT2)
# schematic: cumulative tyre loss crossing the stop cost
px, py, pw, ph = 6.75, 1.95, 5.98, 4.6
rect(s, px, py, pw, ph, PANEL, line=LINE)
label(s, px + 0.3, py + 0.25, "Schematic · the intuition")
ox, oy, ex, ey = px + 0.65, py + ph - 0.7, px + pw - 0.35, py + 0.85
line(s, ox, oy, ex, oy, TEXT3, 1)
line(s, ox, oy, ox, ey, TEXT3, 1)
f = lambda t: 0.06 * t + 0.94 * t ** 2.3
pts = [(ox + t / 40 * (ex - ox), oy - f(t / 40) * (oy - ey)) for t in range(41)]
poly(s, pts, RED, 3)
yp = oy - 0.5 * (oy - ey)
line(s, ox, yp, ex, yp, TEXT, 1.5, dash=True)
tb(s, ex - 2.6, yp - 0.34, 2.6, 0.3, "PIT STOP · 22.6 S", size=12, font=MONO, color=TEXT, align="r")
tc = next(t / 400 for t in range(401) if f(t / 400) >= 0.5)
xc = ox + tc * (ex - ox)
line(s, xc, ey - 0.1, xc, oy, RED, 1.25, dash=True)
dot(s, xc, yp, 0.2, RED, TEXT)
tag(s, xc - 0.42, ey - 0.5, 0.84, 0.36, "Box", size=15)
tb(s, ox + 0.15, oy - 1.0, 2.6, 0.5, "time lost to worn tyres,\nadded up lap by lap", size=12, color=SOFT)
label(s, ex - 2.4, oy + 0.06, "laps on this set  →", w=2.4, align="r")
tb(s, px + 0.3, py + ph - 0.3, pw - 0.6, 0.25, "The engine solves the whole remaining race, Safety Car odds included.", size=11, color=TEXT3)

# 03 · THE DATA
s = new_slide(3, "The data", note=(
    "Every race from 2021 to 2025 from FastF1 timing, rebuilt offline: 114 races and 125,046 laps. Unlike v1, every lap is kept, "
    "including in-laps, out-laps and Safety Car laps, because those are exactly what a pit-stop model needs. "
    "The split is by whole seasons: fit on 2021 to 2023, tune on 2024, and score 2025 exactly once at the end."))
kicker(s, "The data")
title(s, [("Five seasons. ", TEXT), ("One honest test.", RED)])
for i, (v, l) in enumerate((("114", "races"), ("125,046", "laps"), ("26", "circuits"), ("5", "seasons · 2021-2025"))):
    x = 0.6 + i * 3.05
    if i:
        line(s, x - 0.2, 2.15, x - 0.2, 3.45, LINE, 1)
    tb(s, x, 2.05, 2.85, 1.0, v, size=62, font=COND, bold=True, italic=True)
    label(s, x + 0.04, 3.15, l, TEXT2, w=2.9, size=12)
segs = (("Fit · 2021-2023", 7.2, PANEL2, TEXT), ("Tune · 2024", 2.45, "123453", TEXT), ("Test · 2025", 2.45, RED, TEXT))
x = 0.6
for t, w, c, tc_ in segs:
    tag(s, x, 4.1, w - 0.06, 0.62, t, fill=c, color=tc_, size=17, skew=0.18, line=LINE if c == PANEL2 else None)
    x += w
for (t, x, w) in (("Model parameters learned only from these three seasons.", 0.75, 6.8),
                  ("Every knob, threshold and prior chosen here.", 7.85, 2.3),
                  ("Touched once, after everything was frozen.", 10.3, 2.3)):
    tb(s, x, 4.9, w, 0.7, t, size=13, color=TEXT2, line=1.15)
rect(s, 0.6, 5.85, 12.13, 0.72, PANEL, line=LINE)
rect(s, 0.6, 5.85, 0.06, 0.72, RED)
tb(s, 0.9, 5.98, 11.6, 0.5, [("Every lap kept: ", {"bold": True, "color": TEXT}),
                             ("in-laps, out-laps and Safety Car laps are exactly what a pit-stop model needs. Representative laps (90,709) are flagged for fitting pace.", {})],
   size=14, color=TEXT2, anchor="m")

# 04 · ACT I
s = new_slide(4, "Act I · v1", note=(
    "Version 1, July to September 2026: a Ridge regression on absolute lap time scored R squared 0.91 on the 2024 dev season. "
    "It looked great. Then the ablation: remove circuit and driver and it collapses to 0.02. It had learned that Monaco laps are slow, "
    "not how tyres wear. The tyre-age coefficient even came out backwards. Along the way the data had duplicated races, which had "
    "inflated an earlier score to 0.948, and a round_number column corrupted in 82 percent of training rows."))
kicker(s, "Act I  ·  v1  ·  July to September 2026")
title(s, [("It looked ", TEXT), ("great.", RED)])
rect(s, 0.6, 1.95, 5.6, 3.55, PANEL, line=LINE)
label(s, 0.95, 2.2, "Ridge regression on lap time · 71 features", TEXT2, w=5)
tb(s, 0.9, 2.45, 5.2, 2.0, [("R² ", {"size": 48, "color": TEXT3}), ("0.91", {})], size=150, font=COND, bold=True, italic=True, line=0.85)
tb(s, 0.95, 4.65, 5.0, 0.6, "on the 2024 dev season. RMSE about 3.5 s.", size=14, color=TEXT2)
line(s, 6.4, 3.72, 7.05, 3.72, RED, 2.5, head=True)
rect(s, 7.13, 1.95, 5.6, 3.55, "1A0D0F", line=RED, lw=1.25)
label(s, 7.48, 2.2, "Same model, circuit and driver removed", SOFT, w=5)
tb(s, 7.43, 2.45, 5.2, 2.0, [("R² ", {"size": 48, "color": TEXT3}), ("0.02", {"color": RED})], size=150, font=COND, bold=True, italic=True, line=0.85)
tb(s, 7.48, 4.65, 5.0, 0.7, "It knew Monaco laps are slow. It knew nothing about tyres.", size=14, color=TEXT)
for i, (h, b) in enumerate((("Duplicated races", "inflated an earlier R² to 0.948"),
                            ("round_number corrupted", "in 82% of training rows"),
                            ("Tyre age backwards", "−0.0045 s per lap older")) ):
    x = 0.6 + i * 4.08
    rect(s, x, 5.75, 3.97, 0.88, PANEL, line=LINE)
    rect(s, x, 5.75, 0.06, 0.88, RED)
    tb(s, x + 0.25, 5.84, 3.6, 0.35, h.upper(), size=15, font=COND, bold=True, italic=True)
    tb(s, x + 0.25, 6.2, 3.6, 0.35, b, size=12.5, color=TEXT2)

# 05 · THE REBUILD
s = new_slide(5, "Act II · the rebuild", note=(
    "The rebuild starts with one idea: give every driver in every race their own base pace. Then the model never needs tyre or fuel "
    "terms to explain why one car or one track is faster, and tyre wear and fuel are learned only from changes within a driver's race. "
    "Every other term has a physical meaning: compound offsets shrunk from global to circuit to race, fuel, a piecewise-linear wear curve, "
    "temperature times age, and traffic. It is one penalised least-squares fit with about 2,500 parameters, solved in closed form."))
kicker(s, "Act II  ·  the rebuild")
title(s, [("Give every driver ", TEXT), ("their own pace.", RED)])
tb(s, 0.6, 1.8, 11.5, 0.7, "Each driver in each race gets a base pace, so tyres and fuel are learned only from changes inside that race. "
   "That was the step v1 was missing.", size=15, color=TEXT2)
cards = (("Base pace", "per driver, per race", "free", "car + driver + track today", TEXT),
         ("Compound", "soft / medium / hard", "3 levels", "global → circuit → race", RED),
         ("Fuel", "per kg on board", "0.0333 s", "per kg per 90 s lap", BLUE),
         ("Tyre wear", "by tyre age", "6 knots", "0, 4, 10, 18, 28, 40 laps", YELLOW),
         ("Temperature", "track temp × age", "faster wear", "when the track is hotter", PURPLE),
         ("Traffic", "car ahead within 1 s", "+0.41 s", "per lap; +0.20 s at 1-2 s", GREEN))
for i, (n, sub, v, cap, c) in enumerate(cards):
    x = 0.6 + i * 2.04
    rect(s, x, 2.6, 1.92, 2.45, PANEL, line=LINE)
    rect(s, x, 2.6, 1.92, 0.06, c)
    tb(s, x + 0.18, 2.82, 1.7, 0.3, f"{i + 1:02d}", size=10.5, font=MONO, color=TEXT3)
    tb(s, x + 0.18, 3.08, 1.7, 0.4, n.upper(), size=20, font=COND, bold=True, italic=True)
    tb(s, x + 0.18, 3.47, 1.7, 0.3, sub, size=11.5, color=TEXT2)
    tb(s, x + 0.18, 3.95, 1.7, 0.5, v, size=26, font=COND, bold=True, italic=True, color=c if c != TEXT else TEXT)
    tb(s, x + 0.18, 4.47, 1.65, 0.5, cap, size=11, color=TEXT3, line=1.05)
rect(s, 0.6, 5.3, 12.13, 0.62, "101016", line=LINE)
tb(s, 0.85, 5.3, 11.7, 0.62, [("lap time", {"color": TEXT}), (" = base[driver, race] + L · ( compound + 0.0333·fuel + wear(age) + age·(temp + circuit + race) + traffic )", {})],
   size=11.5, font=MONO, color=TEXT2, anchor="m")
tb(s, 0.6, 6.12, 12.1, 0.6, "One penalised least-squares fit, about 2,500 parameters, solved in closed form. Gaussian priors shrink thin data toward the circuit, "
   "then the global value. L scales everything to a 90 s lap. Lap-to-lap noise: 0.41 s.", size=12, color=TEXT3, line=1.1)

# 06 · THE TRICK
s = new_slide(6, "The trick", note=(
    "Inside one stint, fuel falls and tyre age rises together, lap for lap, so you cannot tell them apart. "
    "At a pit stop the tyres reset to zero but the fuel keeps falling. That break is where the model separates the two. "
    "The learned fuel effect, 0.0333 seconds per kilogram, about 0.058 seconds per lap, sits inside the range quoted for current F1 cars, "
    "which is the main sanity check that the separation worked."))
kicker(s, "The trick")
title(s, [("The pit stop splits ", TEXT), ("fuel from wear.", RED)])
tb(s, 0.6, 2.0, 5.3, 2.2, "Inside one stint, fuel falls and tyre age rises together, lap for lap. You cannot tell them apart.\n\n"
   "At a pit stop the tyres reset to zero but the fuel keeps falling. That break is where the model learns each effect.",
   size=15.5, color=TEXT2, line=1.2)
tb(s, 0.6, 4.35, 5.3, 1.0, [("0.0333", {}), (" s/kg", {"size": 28, "color": TEXT2})], size=68, font=COND, bold=True, italic=True, color=BLUE)
tb(s, 0.62, 5.4, 5.0, 0.8, "About 0.058 s per lap of fuel burned: inside the range quoted for current F1 cars.", size=13, color=TEXT3)
px, py, pw, ph = 6.5, 1.95, 6.23, 4.7
rect(s, px, py, pw, ph, PANEL, line=LINE)
L0, L1 = px + 0.55, px + pw - 0.35
X = lambda lap: L0 + (lap - 1) / 56 * (L1 - L0)
top0, top1 = py + 0.75, py + 2.2
bot0, bot1 = py + 2.75, py + 4.1
label(s, px + 0.3, py + 0.3, "Fuel load", BLUE)
label(s, px + 0.3, py + 2.35, "Tyre age", SOFT)
line(s, L0, top1, L1, top1, LINE, 1)
line(s, L0, bot1, L1, bot1, LINE, 1)
poly(s, [(X(1), top0), (X(57), top1)], BLUE, 3)
stints = ((1, 20), (21, 40), (41, 57))
for a, b in stints:
    poly(s, [(X(a), bot1), (X(b), bot1 - (b - a + 1) / 22 * (bot1 - bot0))], SOFT, 3)
    if a > 1:
        line(s, X(a - 0.5), top0 - 0.15, X(a - 0.5), bot1, TEXT3, 1, dash=True)
        tag(s, X(a - 0.5) - 0.32, py + 0.32, 0.64, 0.3, "Pit", size=12, fill=PANEL2, line=LINE)
tb(s, X(21) + 0.08, bot1 + 0.1, 2.4, 0.3, "tyres reset to zero", size=12, color=SOFT)
tb(s, X(23), top0 + (22 / 56) * (top1 - top0) - 0.42, 2.6, 0.3, "fuel keeps falling", size=12, color=BLUE)
label(s, L1 - 2.0, bot1 + 0.12, "lap 1 → 57", w=2.0, align="r")

# 07 · TYRES
s = new_slide(7, "What the tyres do", note=(
    "These are the learned wear curves, fuel removed, as time lost against a fresh set of the same compound. "
    "At 30 laps old a soft costs about 1.7 seconds a lap, a medium about 2.0 and a hard about 1.6. "
    "Running within a second of the car ahead costs another 0.41 seconds a lap. Beyond the 95th-percentile stint length at each circuit "
    "the optimiser adds a cliff penalty rather than trusting the extrapolation."))
kicker(s, "What the tyres do")
title(s, [("Wear, ", TEXT), ("measured.", RED)])
tb(s, 0.6, 1.78, 7.2, 0.5, "Time lost per lap against a fresh set of the same compound, fuel removed. Fitted on all seasons.", size=14, color=TEXT2)
cd = CategoryChartData()
cd.categories = ["new", "10 laps", "20 laps", "30 laps"]
for nme, vals in (("Soft", (0, 0.255, 0.867, 1.718)), ("Medium", (0, 0.728, 1.337, 1.996)), ("Hard", (0, 0.551, 1.076, 1.589))):
    cd.add_series(nme, vals)
rect(s, 0.6, 2.4, 7.4, 4.25, PANEL, line=LINE)
ch = s.shapes.add_chart(XL_CHART_TYPE.LINE_MARKERS, Inches(0.75), Inches(2.55), Inches(7.1), Inches(4.0), cd).chart
chart_clean(ch)
ch.has_title = False
ch.has_legend = True
ch.legend.position, ch.legend.include_in_layout = XL_LEGEND_POSITION.TOP, False
ch.legend.font.size, ch.legend.font.name = Pt(13), COND
ch.legend.font.bold = True
ch.legend.font.color.rgb = rgb(TEXT)
for ser, c in zip(ch.plots[0].series, (SOFT, MED, HARD)):
    ser.smooth = False
    ser.format.line.color.rgb = rgb(c)
    ser.format.line.width = Pt(3)
    ser.marker.style, ser.marker.size = XL_MARKER_STYLE.CIRCLE, 9
    ser.marker.format.fill.solid()
    ser.marker.format.fill.fore_color.rgb = rgb(c)
    ser.marker.format.line.color.rgb = rgb(PANEL)
va, ca = ch.value_axis, ch.category_axis
va.minimum_scale, va.maximum_scale, va.major_unit = 0, 2.5, 0.5
va.has_major_gridlines = True
va.major_gridlines.format.line.color.rgb = rgb(HAIR)
va.format.line.fill.background()
va.tick_labels.number_format, va.tick_labels.number_format_is_linked = '+0.0" s";-0.0" s";0', False
va.tick_labels.font.size = Pt(11)
ca.format.line.color.rgb = rgb(LINE)
ca.major_tick_mark = XL_TICK_MARK.NONE
ca.tick_labels.font.size, ca.tick_labels.font.name = Pt(12), COND
ca.tick_labels.font.bold = True
for i, (n, v, c, sub) in enumerate((("Soft", "+1.72 s", SOFT, "per lap at 30 laps old"), ("Medium", "+2.00 s", MED, "per lap at 30 laps old"),
                                      ("Hard", "+1.59 s", HARD, "per lap at 30 laps old"), ("Traffic", "+0.41 s", GREEN, "per lap, within 1 s of a car"))):
    y = 2.4 + i * 1.08
    rect(s, 8.25, y, 4.48, 0.98, PANEL, line=LINE)
    rect(s, 8.25, y, 0.06, 0.98, c)
    label(s, 8.55, y + 0.14, n, c if c != HARD else TEXT)
    tb(s, 8.55, y + 0.38, 2.0, 0.55, v, size=30, font=COND, bold=True, italic=True)
    tb(s, 10.45, y + 0.42, 2.2, 0.5, sub, size=12, color=TEXT2, line=1.0)

# 08 · PIT LOSS
s = new_slide(8, "What a stop costs", note=(
    "Pit loss is measured, not assumed: for every real stop, the in-lap plus out-lap minus what those laps would have taken without stopping. "
    "That is 1.53 seconds of error per stop against 2.16 for a fixed 22 seconds. Under a Safety Car a lap-time comparison is misleading because "
    "the queue gives the time back, so those stops are measured by the gap to the leader three laps later against matched cars that did not stop. "
    "An SC stop costs 9 percent of a green one, a VSC stop 59 percent, but teams behave as if an SC stop costs about half, because track position is hard to win back."))
kicker(s, "What a stop costs")
title(s, [("22.6 seconds. ", TEXT), ("Unless the Safety Car is out.", RED)])
rows = (("Green flag", 1.0, "22.6 s", GREEN), ("Virtual Safety Car", 0.586, "13.3 s", YELLOW), ("Safety Car", 0.086, "1.9 s", YELLOW))
bx, bw = 2.95, 4.6
for i, (n, frac, v, c) in enumerate(rows):
    y = 2.15 + i * 1.02
    label(s, 0.6, y + 0.12, n, TEXT, w=2.3, size=13)
    rect(s, bx, y, bw, 0.56, PANEL, line=LINE)
    rect(s, bx, y, bw * frac, 0.56, c)
    if frac > 0.8:
        tb(s, bx + bw * frac - 1.3, y, 1.15, 0.56, v, size=15, font=MONO, bold=True, anchor="m", align="r", color=BG)
    else:
        tb(s, bx + bw * frac + 0.15, y, 1.2, 0.56, v, size=15, font=MONO, bold=True, anchor="m")
    tb(s, bx, y + 0.6, 4, 0.3, f"{round(frac * 100)}% of a green stop", size=11, color=TEXT3)
rect(s, bx, 2.15 + 2 * 1.02, bw * 0.5, 0.56, None, line=TEXT, lw=1.25, dash=True)
tb(s, bx + bw * 0.5 - 0.1, 2.15 + 2 * 1.02 + 0.6, 2.5, 0.3, "↑ how teams treat it: about 50%", size=11, color=TEXT)
rect(s, 0.6, 5.45, 6.95, 1.2, PANEL, line=LINE)
tb(s, 0.85, 5.55, 6.5, 1.0, [[("Green: ", {"bold": True, "color": TEXT}), ("in-lap + out-lap minus the two laps the model expected. 1.53 s error per stop vs 2.16 s for a fixed 22 s (465 stops).", {})],
                            [("Safety Car: ", {"bold": True, "color": TEXT}), ("gap to the leader three laps later, against cars that did not stop.", {})]],
   size=12.5, color=TEXT2, line=1.1, after=4)
# Safety Car chain
px, py = 7.95, 2.15
rect(s, px, py, 4.78, 4.5, PANEL, line=LINE)
label(s, px + 0.3, py + 0.25, "Safety Car risk · per-circuit Markov chain")
nodes = {"GREEN": (px + 2.39, py + 1.45, GREEN), "SC": (px + 1.1, py + 3.35, YELLOW), "VSC": (px + 3.68, py + 3.35, YELLOW)}
for k, (cx, cy, c) in nodes.items():
    dot(s, cx, cy, 1.05, PANEL2, c)
    tb(s, cx - 0.5, cy - 0.2, 1.0, 0.4, k, size=17, font=COND, bold=True, italic=True, align="c", anchor="m")
line(s, px + 2.0, py + 1.9, px + 1.4, py + 2.85, TEXT2, 1.25, head=True)
line(s, px + 2.78, py + 1.9, px + 3.38, py + 2.85, TEXT2, 1.25, head=True)
tb(s, px + 0.15, py + 1.85, 1.35, 0.6, "1.2% per lap\nmean 4.1 laps", size=11, font=MONO, color=TEXT2, align="r")
tb(s, px + 3.3, py + 1.85, 1.45, 0.6, "0.95% per lap\nmean 2.1 laps", size=11, font=MONO, color=TEXT2)
tb(s, px + 0.3, py + 4.05, 4.2, 0.35, "Shrunk toward the global rate with a 200-lap prior.", size=11, color=TEXT3)

# 09 · THE CALL
s = new_slide(9, "The call", note=(
    "The call comes from a stochastic dynamic programme. At each lap the state is the compound, tyre age, whether two compounds have been used, "
    "and the track status. Working backwards from the flag, it computes the best expected remaining race time from every state, staying out or pitting. "
    "Safety Car states follow the circuit's chain, so it values staying out for a cheap stop. 30 refits of the tyre model on resampled races vote, "
    "which gives the agreement ring. The same optimiser runs in the browser, and a test checks it makes the same calls as Python."))
kicker(s, "The call")
title(s, [("Solve the rest of the race ", TEXT), ("backwards.", RED)])
rect(s, 0.6, 2.0, 6.5, 0.85, PANEL2, line=LINE)
tb(s, 0.85, 2.0, 6.1, 0.85, [("STATE  ", {"color": RED, "font": COND, "bold": True, "size": 15}), ("lap 15 · soft · 14 laps old · green", {})],
   size=13.5, font=MONO, anchor="m")
for i, (n, f_, c) in enumerate((("Stay out", "lap on worn tyres\n+ best future from age + 1", TEXT), ("Box", "lap + pit loss(status)\n+ best future on a fresh set", RED))):
    x = 0.6 + i * 3.32
    rect(s, x, 3.2, 3.18, 1.35, PANEL, line=c if c == RED else LINE, lw=1.25 if c == RED else 0.75)
    tb(s, x + 0.25, 3.3, 2.8, 0.45, n.upper(), size=22, font=COND, bold=True, italic=True, color=c)
    tb(s, x + 0.25, 3.78, 2.8, 0.7, f_, size=12, font=MONO, color=TEXT2, line=1.1)
line(s, 2.2, 2.85, 2.2, 3.2, TEXT3, 1.25, head=True)
line(s, 5.5, 2.85, 5.5, 3.2, TEXT3, 1.25, head=True)
line(s, 2.2, 4.55, 2.2, 4.9, TEXT3, 1.25, head=True)
line(s, 5.5, 4.55, 5.5, 4.9, TEXT3, 1.25, head=True)
rect(s, 0.6, 4.9, 6.5, 0.7, "1A0D0F", line=RED)
tb(s, 0.85, 4.9, 6.1, 0.7, [("CALL  ", {"color": RED, "font": COND, "bold": True, "size": 15}), ("the cheaper branch · then 30 refits vote", {})], size=13.5, font=MONO, anchor="m")
tb(s, 0.6, 5.8, 6.5, 0.9, "Safety Car states follow the circuit's chain, so staying out for a cheap stop has value. No finish without two compounds. "
   "Past the 95th-percentile stint, a cliff penalty instead of extrapolation.", size=12, color=TEXT3, line=1.1)
pic(s, "garage.jpg", 7.45, 2.0, 5.28, 2.11)
pic(s, "call.jpg", 7.45, 4.2, 5.28, 1.23)
tb(s, 7.45, 5.6, 5.28, 1.0, "The Pit Wall view: the call as team radio, 30-model agreement, the time it is worth. The browser runs a line-for-line port of the "
   "optimiser, tested against Python.", size=12, color=TEXT3, line=1.1)

# 10 · RESULTS
s = new_slide(10, "Results", note=(
    "Results on 2025, a season the model never saw. During a race, the average lap-time error is 0.616 seconds, against 1.111 for simply repeating "
    "the last lap: 45 percent less error. Before the race it is 1.556 seconds, better than the v1 Ridge. Pit loss per stop: 1.53 seconds of error. "
    "Remaining race time from lap 10: within 0.49 percent."))
kicker(s, "Results  ·  2025, never seen in fitting")
title(s, [("45% less error ", TEXT), ("than the obvious guess.", RED)])
rect(s, 0.6, 1.95, 7.0, 4.7, PANEL, line=LINE)
label(s, 0.9, 2.15, "Lap-time error during a race · seconds (MAE)", TEXT2, w=6.5)
hbar_chart(s, 0.75, 2.45, 6.7, 3.75, ["Pitwall", "Repeat the last lap", "Mean of last 5 laps", "Linear fit on the race"],
           [0.616, 1.111, 1.148, 2.549], [RED, "4A4A58", "4A4A58", "4A4A58"], '0.000" s"', 3.1)
tb(s, 0.9, 6.25, 6.6, 0.3, "109,883 predictions across the 2025 season", size=11, color=TEXT3)
tiles = (("0.446 s", "1-5 laps ahead", "error forecasting the next few laps"), ("1.556 s", "before the race", "v1 Ridge: 1.708 s · R² 0.955"),
         ("1.53 s", "pit loss per stop", "fixed 22 s: 2.16 s · 465 stops"), ("0.49%", "rest-of-race time", "from lap 10 · naive: 1.32%"))
for i, (v, l, sub) in enumerate(tiles):
    x, y = 7.85 + (i % 2) * 2.47, 1.95 + (i // 2) * 2.4
    rect(s, x, y, 2.38, 2.3, PANEL, line=LINE)
    tb(s, x + 0.22, y + 0.25, 2.1, 0.8, v, size=38, font=COND, bold=True, italic=True, color=PURPLE if i == 0 else TEXT)
    label(s, x + 0.24, y + 1.15, l, TEXT, w=2.0, size=12)
    tb(s, x + 0.24, y + 1.48, 1.95, 0.7, sub, size=11.5, color=TEXT3, line=1.05)

# 11 · THE 0.81 TRAP
s = new_slide(11, "The 0.81 trap", note=(
    "A published Bi-LSTM pit predictor reports F1 0.81. Pitwall can match it: give the classifier the lap's own time and it scores 0.84, ROC-AUC 0.994. "
    "But a lap that ends in the pit lane is slow because of the stop, so the model is reading the answer off the lap. Using only what is known before "
    "the lap, the honest pit-call model scores F1 0.27, ROC-AUC 0.876. We had our own version of this leak: end-of-lap gap features lifted F1 to 0.32 "
    "until they were moved to lap n minus 1."))
kicker(s, "The 0.81 trap")
title(s, [("The best-looking score ", TEXT), ("was a leak.", RED)])
rect(s, 0.6, 1.95, 7.0, 4.7, PANEL, line=LINE)
label(s, 0.9, 2.15, "Predicting real pit stops, per lap · F1 score", TEXT2, w=6.5)
hbar_chart(s, 0.75, 2.45, 6.7, 3.75, ["Pitwall, shown the lap's own time", "Published Bi-LSTM", "Pitwall pit-call model", "Optimiser alone"],
           [0.839, 0.81, 0.273, 0.159], [RED, "4A4A58", TEXT, "4A4A58"], "0.00", 1.0)
tb(s, 0.9, 6.25, 6.6, 0.3, "2025 season · 22,331 lap states · 625 real stops", size=11, color=TEXT3)
tb(s, 7.95, 2.0, 4.78, 2.4, [[("A lap that ends in the pit lane is slow ", {}), ("because", {"italic": True}), (" of the stop.", {})],
                             [("Give a model that lap's time and it reads the answer off the lap.", {"color": TEXT2, "size": 16})]],
   size=24, font=COND, bold=True, line=1.05, after=10)
for i, (k, v, c) in enumerate((("Leaky", "F1 0.84 · ROC-AUC 0.994", RED), ("Honest", "F1 0.27 · ROC-AUC 0.876", TEXT))):
    y = 4.45 + i * 0.62
    tag(s, 7.95, y, 1.15, 0.42, k, fill=c if c == RED else PANEL2, color=TEXT, line=None if c == RED else LINE)
    tb(s, 9.3, y, 3.4, 0.42, v, size=13.5, font=MONO, anchor="m")
tb(s, 7.95, 5.75, 4.78, 0.9, "Our own leak: end-of-lap gap features lifted F1 from 0.27 to 0.32 until they were moved to lap n−1.", size=12.5, color=TEXT3, line=1.1)


# 12-13 · DNF boards
def dnf(n, section, ttl, rows, note):
    s = new_slide(n, section, note=note)
    kicker(s, section)
    title(s, ttl)
    for j, h in enumerate(("#", "Symptom", "Cause", "Fix")):
        label(s, (0.65, 1.38, 5.05, 8.75)[j], 1.85, h)
    for i, (sym, cause, fix) in enumerate(rows):
        y = 2.12 + i * 0.63
        rect(s, 0.6, y, 12.13, 0.57, PANEL, line=LINE)
        tag(s, 0.68, y + 0.12, 0.55, 0.33, f"{i + 1}", size=13, fill=RED)
        tb(s, 1.38, y, 3.55, 0.57, sym.upper(), size=14.5, font=COND, bold=True, italic=True, anchor="m", line=0.95)
        tb(s, 5.05, y, 3.55, 0.57, cause, size=11.5, color=TEXT2, anchor="m", line=1.0)
        tb(s, 8.75, y, 2.95, 0.57, fix, size=11.5, color=TEXT, anchor="m", line=1.0)
        tag(s, 11.82, y + 0.13, 0.82, 0.31, "Fixed", size=11, fill="12301B", color="7BE39A")
    return s


dnf(12, "What broke  ·  the model", [("Did not ", TEXT), ("finish.", RED)], (
    ("R² 0.948 that would not reproduce", "popular races counted several times in the data", "deduplicated; honest baseline R² 0.91"),
    ("Races mixed together", "round_number corrupted in 82% of training rows", "race IDs rebuilt from file names"),
    ("No pit laps to learn from", "FastF1 IsAccurate drops in-laps, out-laps and SC laps", "keep every lap; flag representative ones"),
    ("Silent garbage in the fit", "groupby returned index labels, not positions", "factorise + argsort split"),
    ("Solver crashed mid-fit", "SciPy positive-definite solve segfaulted", "numpy.linalg.solve"),
    ("Pit loss nonsense under SC", "the Safety Car queue gives the time back", "gap to leader vs matched non-stoppers"),
    ("F1 0.32, too good", "end-of-lap gap leaked the in-lap", "features from lap n−1: honest 0.27")),
    "Here is what broke while building the model, as a DNF board. Each one cost time; each one changed how the project works. "
    "The most important: the pit and Safety Car laps were being thrown away by a FastF1 filter, and a groupby returned labels instead of positions, "
    "which silently corrupted the design matrix without raising any error.")
dnf(13, "What broke  ·  the app", [("And on ", TEXT), ("the track.", RED)], (
    ("3D view crashed with NaN", "service worker served a stale tracks.json", "network-first for data; cache-first only for hashed files"),
    ("Race replay froze for good", "one exception inside requestAnimationFrame", "schedule the next frame first"),
    ("Chase cam 100s of metres behind", "position lerp at 64× playback", "rig derived from the car's own pose"),
    ("Car flipped in corners", "rotation degenerate near a 180° heading", "roll-free yaw + pitch poses"),
    ("Camera wobble at 1×", "telemetry jitter, arc-length error", "smoothed centreline, damped shared heading"),
    ("Front-wing label never drew", "one of seven Html roots failed to mount", "one projected DOM overlay"),
    ("Chicanes missing", "smoothing flattened the geometry", "corners from braking zones in the speed trace")),
    "The app broke in its own ways. A stale cached file crashed the 3D view, one exception could freeze the replay forever, the car flipped "
    "in corners because a rotation was degenerate near a 180 degree heading, and smoothing the track to stop the wobble also erased chicanes, "
    "so corners are now found from braking zones in the real speed trace.")

# 14 · EVOLUTION
s = new_slide(14, "How it evolved", note=(
    "The timeline. v1 in July 2026 built the data pipeline and the Ridge baseline. September brought per-compound models and mostly negative results, "
    "plus the round_number fix. On 2026-10-06 v2 rebuilt everything around the structural model, the measured pit loss, the Safety Car chain and the optimiser. "
    "v2.1 added recency weighting, which cut in-race error by 18 percent, and the leak-free pit-call model. v2.2 turned the dashboard into a cinematic 3D app, "
    "and v2.3 is the F1 broadcast redesign with the What If engine, onboard telemetry, the race-engineer guide and the corner cinematics."))
kicker(s, "How it evolved")
title(s, [("From notebook ", TEXT), ("to pit wall.", RED)])
steps = (("2026-07-10", "v1", "FastF1 pull, dedup and splits; Ridge baseline on lap time", False),
         ("2026-09-03", "v1.x", "per-compound models; tyre effect barely visible; round_number fix", False),
         ("2026-10-06", "v2", "structural model, measured pit loss, SC chain, DP optimiser", True),
         ("2026-10-06", "v2.1", "recency-weighted in-race fit (−18% error); leak-free pit-call model", True),
         ("2026-10-06", "v2.2", "cinematic 3D app: real car model, race theatre", True),
         ("2026-10-06", "v2.3", "F1 broadcast redesign, What If, onboard, guide, corner cinematics", True))
ly = 3.62
line(s, 0.75, ly, 12.6, ly, LINE, 2)
line(s, 0.75 + 2 * 2.04, ly, 12.6, ly, RED, 3)
for i, (d, v, cap, new) in enumerate(steps):
    x = 0.6 + i * 2.04
    tb(s, x, 2.08, 1.95, 0.3, d, size=11, font=MONO, color=TEXT3)
    tb(s, x, 2.38, 1.95, 0.9, v.upper(), size=44, font=COND, bold=True, italic=True, color=TEXT if new else TEXT2)
    dot(s, x + 0.17, ly, 0.3, RED if new else PANEL2, TEXT if new else TEXT3)
    tb(s, x, 4.0, 1.88, 1.6, cap, size=12.5, color=TEXT2 if new else TEXT3, line=1.12)
for i, t in enumerate(("ablate first", "negatives count", "model the base", "leaks hide in time", "make it watchable", "make it stable")):
    tag(s, 0.6 + i * 2.04, 5.2, 1.88, 0.38, t, fill=PANEL2 if i < 2 else "2A0B0A", color=TEXT2 if i < 2 else TEXT, size=12, line=LINE if i < 2 else RED)
label(s, 0.6, 4.88, "What each step taught", TEXT3, w=4)
rect(s, 0.6, 5.85, 12.13, 0.72, PANEL, line=LINE)
rect(s, 0.6, 5.85, 0.06, 0.72, RED)
tb(s, 0.9, 5.85, 11.7, 0.72, "Same data from July to October. What changed was the question asked of it: not “how fast is this lap?” but "
   "“how much slower than this driver's own pace today, and why?”", size=14, color=TEXT, anchor="m")

# 15 · THE PRODUCT
s = new_slide(15, "The product", note=(
    "The model became an app you can watch. Six views: Pit Wall for the call, Race Theatre replaying every 2025 race in 3D on the real circuits with all 20 cars "
    "at their recorded pace, Onboard with real telemetry, What If, Tyre Lab and Proof. A personal race-engineer guide tours it at your level, "
    "and Explain mode explains anything you click."))
kicker(s, "The product")
title(s, [("Six views. ", TEXT), ("One race engineer.", RED)])
pic(s, "theatre_shot.jpg", 0.6, 1.95, 6.4, 4.6)
pic(s, "garage.jpg", 7.2, 1.95, 5.53, 2.21)
pic(s, "degradation_app.jpg", 7.2, 4.34, 5.53, 2.21)
for (x, y, t) in ((0.78, 6.07, "Race Theatre · Director camera"), (7.38, 3.68, "Pit Wall · spec hotspots"), (7.38, 6.07, "Tyre Lab · live curves")):
    tag(s, x, y, 2.95 if "Theatre" in t else 2.3, 0.34, t, fill=BG, color=TEXT, size=11.5, line=LINE)
tb(s, 0.6, 6.62, 12.1, 0.3, [("Pit Wall · Race Theatre · Onboard · What If · Tyre Lab · Proof", {"color": TEXT2}),
                             ("     live: b25bb1004-wq.github.io/F1-Pitstop_Strategy", {"color": RED, "font": MONO, "size": 11})], size=12, font=COND, bold=True)

# 16 · WHAT IF
s = new_slide(16, "What If", note=(
    "What If reruns a real 2025 race with one thing changed. You ask in plain English; the parser shows what it understood and reports anything it cannot place. "
    "Each driver's real clean pace is taken lap by lap, the change goes through the model, and Safety Cars are queue events where lapped cars keep their deficit. "
    "Each what-if clock is the real clock plus the modelled difference, so an unchanged race reproduces reality exactly. "
    "Here, if Leclerc had followed Pitwall's strategy at Monza, he finishes second instead of fourth, 13 seconds quicker."))
kicker(s, "What If")
title(s, [("Change one thing. ", TEXT), ("Rerun the race.", RED)])
rect(s, 0.6, 1.95, 7.9, 0.58, PANEL2, line=LINE)
rect(s, 0.6, 1.95, 0.06, 0.58, RED)
tb(s, 0.88, 1.95, 6.0, 0.58, "What if Leclerc had followed Pitwall's strategy at Monza?", size=14, color=TEXT, anchor="m")
tag(s, 7.15, 2.04, 1.22, 0.4, "Run it", size=13)
pic(s, "whatif.jpg", 0.6, 2.7, 7.9, 2.82)
for i, (k, v) in enumerate((("Result", "P4 → P2"), ("Race time", "−13.0 s"), ("Strategy", "lap 28 for softs"))):
    x = 0.6 + i * 2.66
    rect(s, x, 5.7, 2.58, 0.92, PANEL, line=LINE)
    label(s, x + 0.2, 5.8, k)
    tb(s, x + 0.2, 6.05, 2.3, 0.5, v, size=24, font=COND, bold=True, italic=True, color=BLUE if i == 1 else (RED if i == 0 else TEXT))
for i, (h, b) in enumerate((("Parse", "Deterministic. Anything it cannot place is reported, never guessed."),
                            ("Real pace", "Each driver's clean pace, lap by lap, pit-lane time removed."),
                            ("Apply", "Tyre curves, pit loss and the optimiser change the laps that differ."),
                            ("Safety Cars", "Queue events; lapped cars keep their deficit."),
                            ("Anchor", "what-if clock = real clock + (what-if − baseline). No change = reality."))):
    y = 1.95 + i * 0.94
    tb(s, 8.85, y, 0.6, 0.5, f"{i + 1:02d}", size=26, font=COND, bold=True, italic=True, color=RED)
    tb(s, 9.5, y + 0.02, 3.2, 0.32, h.upper(), size=15, font=COND, bold=True, italic=True)
    tb(s, 9.5, y + 0.33, 3.23, 0.58, b, size=11.5, color=TEXT2, line=1.05)

# 17 · CINEMATICS
s = new_slide(17, "Making it feel real", note=(
    "Making it feel real. Corners are found from braking zones in the fastest-lap speed trace: every minimum at least 20 kilometres per hour below its run-up. "
    "At Monza that finds all six real braking zones. A trackside camera sits on the outside of each corner with a long lens that zooms to hold the car's size. "
    "The front wheels steer from the real curvature, and every pose is yaw and pitch only, so the car can never flip. The Director mode cuts like a broadcast."))
kicker(s, "Making it feel real")
title(s, [("Directed like ", TEXT), ("a broadcast.", RED)])
pic(s, "onboard.jpg", 0.6, 1.95, 7.35, 4.04)
tag(s, 0.78, 5.5, 3.3, 0.34, "Onboard · long lens through Monza's Roggia", fill=BG, color=TEXT, size=11.5, line=LINE)
tb(s, 0.6, 6.15, 7.35, 0.5, "Corner caption, apex speed and gear come from the real fastest lap.", size=12, color=TEXT3)
for i, (h, f_, b) in enumerate((("Corners from braking", "Δv ≥ 20 km/h", "Monza: all 6 braking zones found"),
                                 ("Long-lens trackside cams", "fov = 2·atan(7.5 m / d)", "outside each corner, zooming to hold the car"),
                                 ("Steering from curvature", "δ = atan(3.6 m × κ)", "front wheels turn into each corner"),
                                 ("Roll-free poses", "yaw + pitch only", "the car can never flip"),
                                 ("Director cuts", "cut at every corner", "chase, side and aerial on straights; steady at 8×+"))):
    y = 1.95 + i * 0.94
    rect(s, 8.2, y, 4.53, 0.84, PANEL, line=LINE)
    rect(s, 8.2, y, 0.06, 0.84, RED)
    tb(s, 8.45, y + 0.07, 4.2, 0.32, h.upper(), size=15, font=COND, bold=True, italic=True)
    tb(s, 8.45, y + 0.42, 2.1, 0.35, f_, size=10.5, font=MONO, color=RED)
    tb(s, 10.55, y + 0.4, 2.12, 0.42, b, size=10.5, color=TEXT2, line=1.0)

# 18 · LESSONS
s = new_slide(18, "Lessons", note=(
    "Five lessons. A high R squared can be a lookup table, so ablate before you believe it. Separate effects where they decouple: the pit stop resets the tyres but not the fuel. "
    "Leaks hide in time, so use only what was known before the lap. Hold out whole seasons and score the test year once. And stability is a feature: smooth inputs and damped rigs beat clever cameras."))
kicker(s, "Lessons")
title(s, [("Five things ", TEXT), ("this taught me.", RED)])
for i, (h, b) in enumerate((("A high R² can be a lookup table.", "Ablate before you believe it: 0.91 became 0.02."),
                            ("Separate effects where they decouple.", "The pit stop resets the tyres, not the fuel."),
                            ("Leaks hide in time.", "Use only what was known before the lap."),
                            ("Hold out whole seasons.", "Tune on 2024, score 2025 once."),
                            ("Stability is a feature.", "Smooth inputs and damped rigs beat clever cameras."))):
    y = 1.98 + i * 0.93
    if i:
        line(s, 0.6, y - 0.06, 12.73, y - 0.06, HAIR, 1)
    tb(s, 0.6, y, 1.0, 0.8, f"{i + 1:02d}", size=40, font=COND, bold=True, italic=True, color=RED, anchor="m")
    tb(s, 1.65, y, 6.4, 0.8, h.upper(), size=27, font=COND, bold=True, italic=True, anchor="m")
    tb(s, 8.2, y, 4.53, 0.8, b, size=14.5, color=TEXT2, anchor="m", line=1.05)

# 19 · NEXT
s = new_slide(19, "Next race", note=(
    "What's next. A two-car undercut and overcut model to close the gap between the time-minimising optimiser and what teams actually do. "
    "Mapping soft, medium and hard to Pirelli's C1 to C5 per event. And FP2 long runs to counter survivorship bias in the tyre curves."))
kicker(s, "Next race")
title(s, [("What's ", TEXT), ("next.", RED)])
for i, (h, b, c) in enumerate((("Two-car undercut model", "The optimiser minimises one car's own time. Undercuts and covering rivals are why real calls differ. The gap features already exist.", RED),
                               ("C1-C5 allocation", "Soft, medium and hard are relative per event. Mapping them to Pirelli's compounds would separate the curves.", YELLOW),
                               ("FP2 long runs", "Very old tyres are mostly seen on cars where they held up. Practice long runs counter that survivorship bias.", BLUE))):
    x = 0.6 + i * 4.1
    rect(s, x, 2.05, 3.93, 3.4, PANEL, line=LINE)
    rect(s, x, 2.05, 3.93, 0.06, c)
    tb(s, x + 0.3, 2.35, 1.0, 0.6, f"{i + 1:02d}", size=30, font=COND, bold=True, italic=True, color=c if c != YELLOW else MED)
    tb(s, x + 0.3, 3.0, 3.4, 0.9, h.upper(), size=25, font=COND, bold=True, italic=True, line=0.95)
    tb(s, x + 0.3, 3.7, 3.35, 1.6, b, size=13.5, color=TEXT2, line=1.15)
rect(s, 0.6, 5.7, 12.13, 0.85, PANEL, line=LINE)
tb(s, 0.9, 5.7, 11.7, 0.85, [("Known limits today: ", {"bold": True, "color": TEXT}),
                             ("What If assumes time gained becomes positions (traffic and overtaking are not modelled). Wet races are excluded. "
                              "Before a race, a new season's car pace is the main unknown.", {})], size=13, color=TEXT2, anchor="m", line=1.1)

# 20 · CHEQUERED FLAG
s = new_slide(20, "Chequered flag", bg="bg_flag.jpg", note=(
    "That's Pitwall: a structural model that knows why a lap is slow, a decision engine that solves the rest of the race, an honest evaluation, "
    "and an app that lets you watch it and ask it what if. The live app and all the code are linked here."))
sq = 0.3
for r in range(2):
    for c in range(int(13.333 / sq) + 1):
        if (r + c) % 2 == 0:
            rect(s, c * sq, 0.0 + r * sq, sq, sq, "E8E8EC")
kicker(s, "Chequered flag", y=1.25)
tb(s, 0.55, 1.6, 12.2, 2.2, [("BOX, ", {}), ("BOX.", {"color": RED})], size=150, font=COND, bold=True, italic=True, line=0.85)
for i, (k, v) in enumerate((("Live app", "b25bb1004-wq.github.io/F1-Pitstop_Strategy"), ("Code", "github.com/b25bb1004-wq/F1-Pitstop_Strategy"),
                            ("How the model works", "notes/how_the_model_works.md"))):
    y = 4.25 + i * 0.55
    label(s, 0.62, y + 0.06, k, TEXT3, w=2.6, size=12)
    tb(s, 3.2, y, 9.0, 0.4, v, size=17, font=MONO, color=TEXT if i < 2 else TEXT2)
tb(s, 0.6, 6.05, 12.1, 0.6, "Arnav Yadav · FastF1 timing data 2021-2025 · frontend built with Claude Code · car model “Ferrari F1-75” by Sketcher, "
   "CC-BY-NC-4.0, de-branded · fonts: Barlow, JetBrains Mono (OFL)", size=11, color=TEXT3, line=1.1)

prs.save(OUT)
print("saved", OUT, len(prs.slides), "slides")
